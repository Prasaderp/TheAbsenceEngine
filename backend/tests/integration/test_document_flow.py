import io
import csv
import uuid
import pytest
from unittest.mock import AsyncMock
import openpyxl
import docx


def _csv_bytes() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["col1", "col2"])
    w.writerow(["val1", "val2"])
    return buf.getvalue().encode()


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["h1", "h2"])
    ws.append(["r1", "r2"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Integration test document content.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def mock_storage(monkeypatch):
    monkeypatch.setattr("app.services.document_service.storage.upload", AsyncMock(return_value="mock/key"))
    monkeypatch.setattr("app.services.document_service.storage.delete", AsyncMock())
    monkeypatch.setattr("app.services.document_service.storage.sha256", lambda data: "a" * 64)
    monkeypatch.setattr("app.services.document_service.storage.build_key", lambda *_: "mock/key")


class TestDocumentUpload:
    @pytest.mark.asyncio
    async def test_upload_txt(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("test.txt", b"Hello world document", "text/plain")},
            headers=auth_headers,
        )
        assert response.status_code == 201
        body = response.json()
        assert body["filename"] == "test.txt"
        assert body["mime_type"] == "text/plain"
        assert "id" in body

    @pytest.mark.asyncio
    async def test_upload_csv(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("data.csv", _csv_bytes(), "text/csv")},
            headers=auth_headers,
        )
        assert response.status_code == 201

    @pytest.mark.asyncio
    async def test_upload_xlsx(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("data.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=auth_headers,
        )
        assert response.status_code == 201

    @pytest.mark.asyncio
    async def test_upload_docx(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("doc.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=auth_headers,
        )
        assert response.status_code == 201

    @pytest.mark.asyncio
    async def test_upload_empty_file_rejected(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("empty.txt", b"", "text/plain")},
            headers=auth_headers,
        )
        assert response.status_code == 422

    @pytest.mark.asyncio
    async def test_upload_unsupported_mime_rejected(self, async_client, auth_headers):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("image.png", b"fake png data", "image/png")},
            headers=auth_headers,
        )
        assert response.status_code == 422

    @pytest.mark.asyncio
    async def test_upload_requires_auth(self, async_client):
        response = await async_client.post(
            "/api/v1/documents",
            files={"file": ("test.txt", b"content", "text/plain")},
        )
        assert response.status_code in (401, 403)


class TestDocumentList:
    @pytest.mark.asyncio
    async def test_list_empty(self, async_client, auth_headers):
        response = await async_client.get("/api/v1/documents", headers=auth_headers)
        assert response.status_code == 200
        body = response.json()
        assert "data" in body
        assert "meta" in body
        assert body["meta"]["page"] == 1

    @pytest.mark.asyncio
    async def test_list_after_upload(self, async_client, auth_headers):
        await async_client.post(
            "/api/v1/documents",
            files={"file": ("list_test.txt", b"List test content", "text/plain")},
            headers=auth_headers,
        )
        response = await async_client.get("/api/v1/documents", headers=auth_headers)
        assert response.status_code == 200
        assert len(response.json()["data"]) >= 1

    @pytest.mark.asyncio
    async def test_pagination_params(self, async_client, auth_headers):
        response = await async_client.get("/api/v1/documents?page=1&per_page=5", headers=auth_headers)
        assert response.status_code == 200
        meta = response.json()["meta"]
        assert meta["per_page"] == 5


class TestDocumentGet:
    @pytest.mark.asyncio
    async def test_get_existing(self, async_client, auth_headers):
        upload = await async_client.post(
            "/api/v1/documents",
            files={"file": ("get_test.txt", b"Get test content", "text/plain")},
            headers=auth_headers,
        )
        doc_id = upload.json()["id"]
        response = await async_client.get(f"/api/v1/documents/{doc_id}", headers=auth_headers)
        assert response.status_code == 200
        assert response.json()["id"] == doc_id

    @pytest.mark.asyncio
    async def test_get_nonexistent_returns_404(self, async_client, auth_headers):
        response = await async_client.get(f"/api/v1/documents/{uuid.uuid4()}", headers=auth_headers)
        assert response.status_code == 404
        assert "error" in response.json()


class TestDocumentDelete:
    @pytest.mark.asyncio
    async def test_delete_existing(self, async_client, auth_headers):
        upload = await async_client.post(
            "/api/v1/documents",
            files={"file": ("del_test.txt", b"Delete test content", "text/plain")},
            headers=auth_headers,
        )
        doc_id = upload.json()["id"]
        response = await async_client.delete(f"/api/v1/documents/{doc_id}", headers=auth_headers)
        assert response.status_code == 204

    @pytest.mark.asyncio
    async def test_deleted_doc_returns_404(self, async_client, auth_headers):
        upload = await async_client.post(
            "/api/v1/documents",
            files={"file": ("del2_test.txt", b"Another delete test", "text/plain")},
            headers=auth_headers,
        )
        doc_id = upload.json()["id"]
        await async_client.delete(f"/api/v1/documents/{doc_id}", headers=auth_headers)
        response = await async_client.get(f"/api/v1/documents/{doc_id}", headers=auth_headers)
        assert response.status_code == 404

    @pytest.mark.asyncio
    async def test_delete_nonexistent_returns_404(self, async_client, auth_headers):
        response = await async_client.delete(f"/api/v1/documents/{uuid.uuid4()}", headers=auth_headers)
        assert response.status_code == 404

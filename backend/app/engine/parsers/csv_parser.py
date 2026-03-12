import csv
import io
import openpyxl
from app.engine.parsers.base import BaseParser, ParsedDocument, Section
from app.shared.errors import ValidationError


def _rows_to_text(headers: list[str], rows: list[list]) -> tuple[str, list[Section]]:
    header_line = " | ".join(str(h) for h in headers)
    data_lines = [" | ".join(str(c) for c in row) for row in rows]
    text = header_line + "\n" + "\n".join(data_lines)
    section = Section(heading="Table", text=text, level=1)
    return text, [section]


class CSVParser(BaseParser):
    async def parse(self, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            text_data = data.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text_data))
            all_rows = list(reader)
            if len(all_rows) < 2:
                raise ValidationError("CSV must have at least a header row and one data row.")
            headers, *rows = all_rows
            full_text, sections = _rows_to_text(headers, rows)
            return ParsedDocument(text=full_text, sections=sections, metadata={"row_count": len(rows)})
        except ValidationError:
            raise
        except Exception as e:
            raise ValidationError(f"CSV parsing failed: {e}") from e


class XLSXParser(BaseParser):
    async def parse(self, data: bytes, mime_type: str) -> ParsedDocument:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            all_sections: list[Section] = []
            all_texts: list[str] = []
            for sheet in wb.worksheets:
                rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
                if len(rows) < 2:
                    continue
                headers = [str(h) if h is not None else "" for h in rows[0]]
                data_rows = [[str(c) if c is not None else "" for c in r] for r in rows[1:]]
                text, secs = _rows_to_text(headers, data_rows)
                all_texts.append(f"[Sheet: {sheet.title}]\n{text}")
                for s in secs:
                    s.heading = f"{sheet.title} — {s.heading}"
                all_sections.extend(secs)
            if not all_texts:
                raise ValidationError("XLSX file contains no usable data.")
            return ParsedDocument(
                text="\n\n".join(all_texts),
                sections=all_sections,
                metadata={"sheet_count": len(wb.worksheets)},
            )
        except ValidationError:
            raise
        except Exception as e:
            raise ValidationError(f"XLSX parsing failed: {e}") from e
